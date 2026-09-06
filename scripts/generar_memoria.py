"""Genera la memoria del tronco común en formato Word.

    python scripts/generar_memoria.py

Decisión deliberada: la tabla de modelos y parámetros y el texto del system
prompt **se leen del código fuente**, no se transcriben. Así el documento no
puede afirmar una temperatura que la aplicación no aplique. Si alguien cambia
`src/bedrock/models.py`, este documento cambia con él en la siguiente
generación.
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

from src.bedrock.models import (  # noqa: E402
    IMAGE_STYLES,
    TEXT_PROFILES,
    TITAN_EMBEDDINGS,
    TextTask,
)
from src.config import ASSETS_DIR, DOCS_DIR  # noqa: E402
from src.prompts.system_prompts import build_system_prompt  # noqa: E402

AZUL = RGBColor(0x1B, 0x3A, 0x5C)
CORAL = RGBColor(0xF0, 0x55, 0x4F)
GRIS = RGBColor(0x5A, 0x62, 0x6B)
GRAFITO = RGBColor(0x23, 0x28, 0x2D)

AUTOR = "José Ruber Moncayo Navia"
TITULO = "Aplicación de generación de imágenes y edición de contenido con Amazon Bedrock"


# ---------------------------------------------------------------------------
# Utilidades de formato
# ---------------------------------------------------------------------------


def configurar_estilos(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for nivel, tamano, color in [(1, 17, AZUL), (2, 14, AZUL), (3, 12, GRAFITO)]:
        estilo = doc.styles[f"Heading {nivel}"]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tamano)
        estilo.font.color.rgb = color
        estilo.font.bold = True
        estilo.paragraph_format.space_before = Pt(16 if nivel == 1 else 12)
        estilo.paragraph_format.space_after = Pt(6)


def parrafo(doc, texto="", *, negrita=False, cursiva=False, tamano=11,
            color=None, alineacion=None, espacio_despues=None, sangria=None):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = negrita
    run.italic = cursiva
    run.font.size = Pt(tamano)
    if color is not None:
        run.font.color.rgb = color
    if alineacion is not None:
        p.alignment = alineacion
    if espacio_despues is not None:
        p.paragraph_format.space_after = Pt(espacio_despues)
    if sangria is not None:
        p.paragraph_format.left_indent = Cm(sangria)
    return p


def parrafo_rico(doc, fragmentos, *, sangria=None, espacio_despues=None):
    """Párrafo con tramos en negrita: fragmentos es [(texto, negrita), ...]."""
    p = doc.add_paragraph()
    for texto, negrita in fragmentos:
        run = p.add_run(texto)
        run.bold = negrita
        run.font.size = Pt(11)
    if sangria is not None:
        p.paragraph_format.left_indent = Cm(sangria)
    if espacio_despues is not None:
        p.paragraph_format.space_after = Pt(espacio_despues)
    return p


def vineta(doc, texto, nivel=0):
    p = doc.add_paragraph(texto, style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.75 + nivel * 0.6)
    p.paragraph_format.space_after = Pt(4)
    return p


def vineta_rica(doc, fragmentos, nivel=0):
    p = doc.add_paragraph(style="List Bullet")
    for texto, negrita in fragmentos:
        run = p.add_run(texto)
        run.bold = negrita
        run.font.size = Pt(11)
    p.paragraph_format.left_indent = Cm(0.75 + nivel * 0.6)
    p.paragraph_format.space_after = Pt(4)
    return p


def sombrear(celda, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    celda._tc.get_or_add_tcPr().append(shd)


def tabla(doc, cabeceras: list[str], filas: list[list[str]],
          anchos: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(cabeceras))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, texto in enumerate(cabeceras):
        celda = t.rows[0].cells[i]
        celda.text = ""
        run = celda.paragraphs[0].add_run(texto)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(celda, "1B3A5C")

    for fila in filas:
        celdas = t.add_row().cells
        for i, texto in enumerate(fila):
            celdas[i].text = ""
            run = celdas[i].paragraphs[0].add_run(str(texto))
            run.font.size = Pt(9.5)

    if anchos:
        for fila in t.rows:
            for i, ancho in enumerate(anchos):
                fila.cells[i].width = Cm(ancho)
    return t


def caja_destacada(doc, titulo: str, cuerpo: str, color_fondo="FFF4E5"):
    """Recuadro para decisiones y advertencias clave."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    celda = t.rows[0].cells[0]
    sombrear(celda, color_fondo)

    celda.text = ""
    p = celda.paragraphs[0]
    run = p.add_run(titulo)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = GRAFITO

    p2 = celda.add_paragraph()
    run2 = p2.add_run(cuerpo)
    run2.font.size = Pt(10)
    doc.add_paragraph()
    return t


def bloque_codigo(doc, texto: str, tamano=8.5):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    celda = t.rows[0].cells[0]
    sombrear(celda, "F5F5F2")
    celda.text = ""
    for i, linea in enumerate(texto.split("\n")):
        p = celda.paragraphs[0] if i == 0 else celda.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(linea)
        run.font.name = "Consolas"
        run.font.size = Pt(tamano)
    doc.add_paragraph()
    return t


def indice(doc):
    """Inserta un campo TOC que Word rellena al abrir el documento."""
    p = doc.add_paragraph()
    run = p.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    texto = OxmlElement("w:t")
    texto.text = "Pulsa aquí con el botón derecho y elige «Actualizar campos» para generar el índice."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    for elemento in (fld_begin, instr, fld_sep, texto, fld_end):
        run._r.append(elemento)


def pie_de_pagina(doc) -> None:
    seccion = doc.sections[0]
    p = seccion.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{AUTOR} · Caso Práctico Unidad 3 · Generative AI · IEP    ")
    run.font.size = Pt(8)
    run.font.color.rgb = GRIS

    run2 = p.add_run()
    for tipo, texto in [("begin", None), (None, "PAGE"), ("separate", None),
                        (None, "1"), ("end", None)]:
        if tipo:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tipo)
            run2._r.append(fld)
        elif texto == "PAGE":
            instr = OxmlElement("w:instrText")
            instr.text = "PAGE"
            run2._r.append(instr)
        else:
            t = OxmlElement("w:t")
            t.text = texto
            run2._r.append(t)
    run2.font.size = Pt(8)
    run2.font.color.rgb = GRIS


# ---------------------------------------------------------------------------
# Secciones del documento
# ---------------------------------------------------------------------------


def portada(doc):
    for _ in range(4):
        doc.add_paragraph()

    parrafo(doc, "INSTITUTO EUROPEO DE POSGRADO", negrita=True, tamano=12,
            color=GRIS, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    parrafo(doc, "Máster en Inteligencia Artificial Generativa", tamano=11,
            color=GRIS, alineacion=WD_ALIGN_PARAGRAPH.CENTER,
            espacio_despues=40)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITULO)
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = AZUL
    p.paragraph_format.space_after = Pt(10)

    parrafo(doc, "Caso Práctico · Unidad 3 · Generative AI", tamano=13,
            color=GRAFITO, alineacion=WD_ALIGN_PARAGRAPH.CENTER,
            espacio_despues=6)
    parrafo(doc, "Vía A — Aplicación funcional", negrita=True, tamano=12,
            color=CORAL, alineacion=WD_ALIGN_PARAGRAPH.CENTER,
            espacio_despues=50)

    parrafo(doc, "Autor", negrita=True, tamano=10, color=GRIS,
            alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=2)
    parrafo(doc, AUTOR, negrita=True, tamano=14,
            alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=24)

    parrafo(doc, "Aplicación entregada:  Aurora Studio", tamano=10, color=GRIS,
            alineacion=WD_ALIGN_PARAGRAPH.CENTER, espacio_despues=2)
    parrafo(doc, "Modo de ejecución:  simulado (sin cuenta AWS)", tamano=10,
            color=GRIS, alineacion=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()


def seccion_indice(doc):
    doc.add_heading("Índice", level=1)
    indice(doc)
    doc.add_page_break()


def seccion_1(doc):
    doc.add_heading("Introducción y alcance", level=1)

    parrafo(doc,
            "Este documento recoge el diseño de Aurora Studio, una herramienta interna "
            "de generación de imágenes y edición de contenido para una agencia de "
            "marketing y publicidad, construida sobre Amazon Bedrock. Cubre las seis "
            "secciones del tronco común exigido por la guía del trabajo práctico (3.1 a "
            "3.6) y documenta además la ejecución de la Vía A, es decir, la aplicación "
            "funcional que acompaña a esta memoria.")

    parrafo(doc,
            "La aplicación cubre las cuatro funcionalidades que pide el enunciado: "
            "generación de imágenes a partir de texto con selección de estilo y galería; "
            "edición de contenido con cuatro operaciones e historial de versiones; "
            "colaboración con roles, permisos y comentarios; y un marco de ética y "
            "seguridad aplicado en el propio flujo de ejecución, no solo enunciado.")

    doc.add_heading("Un aviso de honestidad sobre el modo de ejecución", level=2)

    parrafo(doc,
            "La aplicación se entrega ejecutándose en modo simulado: no se ha "
            "contratado una cuenta de AWS y, por tanto, ninguna de las respuestas que "
            "muestran las capturas procede realmente de Amazon Bedrock. La guía del "
            "trabajo contempla expresamente esta posibilidad y no la penaliza, pero "
            "exige que la integración esté bien planteada. Esa exigencia es la que ha "
            "gobernado la decisión de arquitectura más importante del proyecto, y se "
            "explica en el apartado 3.2.")

    caja_destacada(
        doc,
        "En una frase",
        "El código que construye las peticiones a Bedrock, las envía y procesa sus "
        "respuestas es código real y se ejecuta íntegro. Lo único simulado es el "
        "transporte de red. Cambiar la variable BEDROCK_BACKEND de «mock» a «aws» "
        "hace que la aplicación llame a AWS de verdad sin modificar ninguna otra "
        "línea.",
        "E8F4EC",
    )


def seccion_31(doc):
    doc.add_heading("3.1. Problema y usuarios", level=1)

    parrafo(doc,
            "La agencia produce piezas creativas —combinaciones de imagen y texto— para "
            "campañas de cliente. El cuello de botella no está en la calidad del equipo, "
            "sino en el número de vueltas que da cada pieza antes de publicarse: el "
            "diseñador necesita propuestas visuales para conversar con el cliente, el "
            "redactor reescribe el mismo texto media docena de veces, y el aprobador "
            "recibe versiones por correo sin saber qué cambió respecto a la anterior.")

    parrafo(doc,
            "Aurora Studio ataca las tres fricciones a la vez: acelera la exploración "
            "visual, asiste la reescritura y conserva la trazabilidad de todo lo que "
            "pasa por la herramienta.")

    doc.add_heading("Historias de usuario", level=2)

    historias = [
        ("Diseñador",
         "generar imágenes desde una descripción de texto y elegir el estilo",
         "acelerar mis propuestas visuales sin esperar a una sesión de fotos"),
        ("Diseñador",
         "reproducir exactamente una imagen que generé ayer usando su semilla",
         "poder iterar sobre una propuesta concreta en lugar de volver a empezar"),
        ("Redactor",
         "resumir, expandir, corregir y generar variaciones de un texto",
         "publicar más rápido y con un estilo más consistente"),
        ("Redactor",
         "que la herramienta aplique la guía de estilo de marca automáticamente",
         "no tener que recordar de memoria las prohibiciones de estilo"),
        ("Aprobador",
         "revisar, comparar y comentar versiones antes de publicar",
         "controlar la calidad y saber exactamente qué cambió entre dos versiones"),
        ("Aprobador",
         "restaurar una versión anterior sin perder el historial",
         "poder rectificar sin destruir la evidencia de lo que se aprobó"),
        ("Responsable legal",
         "conocer el prompt y la semilla de cualquier imagen publicada",
         "poder acreditar cómo se produjo si un tercero la cuestiona"),
    ]

    for rol, accion, beneficio in historias:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.space_after = Pt(6)
        for texto, negrita in [("Como ", False), (rol, True), (", quiero ", False),
                               (accion, True), (", para ", False),
                               (beneficio + ".", False)]:
            run = p.add_run(texto)
            run.bold = negrita
            run.font.size = Pt(11)

    parrafo(doc,
            "Las tres primeras figuras son las que exige el enunciado. La cuarta —el "
            "responsable legal— no aparece en él, pero se ha incorporado porque es "
            "quien realmente impone los requisitos de trazabilidad y de copyright que "
            "estructuran la sección 3.6. Ignorarlo habría dejado esos controles sin "
            "un usuario que los reclamara, y por tanto sin justificación funcional.")


def seccion_32(doc):
    doc.add_heading("3.2. Arquitectura del sistema", level=1)

    parrafo(doc,
            "El sistema se organiza en cinco capas. La pieza central es el wrapper: la "
            "capa propia que recibe la petición del usuario, aplica las reglas de "
            "negocio y de seguridad, decide qué modelo invocar y con qué parámetros, "
            "llama a Bedrock y procesa la respuesta antes de devolverla.")

    diagrama = ASSETS_DIR / "arquitectura.png"
    if diagrama.exists():
        doc.add_picture(str(diagrama), width=Cm(16.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        parrafo(doc, "Figura 1. Arquitectura de Aurora Studio.", tamano=9,
                color=GRIS, alineacion=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_heading("Responsabilidad de cada capa", level=2)

    tabla(doc,
          ["Capa", "Responsabilidad", "Implementación"],
          [
              ["Interfaz", "Cuatro pantallas: generación, edición, colaboración y galería.",
               "Streamlit · src/ui/"],
              ["Wrapper",
               "Moderación, defensa anti-inyección, construcción de prompts, "
               "recuperación RAG, perfiles de inferencia, dominio.",
               "src/services/, src/security/, src/prompts/, src/domain/"],
              ["Transporte",
               "Serializa el cuerpo JSON, invoca invoke_model y parsea la respuesta. "
               "Traduce errores de AWS a mensajes accionables.",
               "src/bedrock/client.py"],
              ["Modelos", "Claude, Stable Diffusion XL y Titan Embeddings.",
               "Amazon Bedrock"],
              ["Almacenamiento",
               "Galería con procedencia, historial de versiones solo-anexado, "
               "índice vectorial de la guía de marca.",
               "Estado de sesión · data/"],
          ],
          anchos=[2.6, 8.0, 5.4])

    doc.add_heading("Justificación: API gestionada frente a modelo propio", level=2)

    parrafo(doc,
            "La alternativa a Bedrock sería desplegar Stable Diffusion y un modelo de "
            "lenguaje abierto en infraestructura propia. Se ha descartado por tres "
            "razones, en este orden de peso:")

    vineta_rica(doc, [
        ("Tiempo de arranque. ", True),
        ("Con Bedrock, la primera llamada funciona el mismo día en que se concede el "
         "acceso al modelo. Con infraestructura propia haría falta aprovisionar "
         "instancias con GPU, instalar los pesos, montar el servicio de inferencia y "
         "resolver el escalado. Para una herramienta interna cuyo valor está en el "
         "flujo de trabajo, no en la inferencia, ese esfuerzo no se recupera.", False),
    ])
    vineta_rica(doc, [
        ("Coste operativo. ", True),
        ("El modelo por uso encaja con un patrón de demanda irregular: los picos "
         "coinciden con los cierres de campaña y hay días de actividad casi nula. Una "
         "GPU reservada se paga esté o no trabajando. La agencia no tiene, además, un "
         "equipo de plataforma que mantenga ese despliegue.", False),
    ])
    vineta_rica(doc, [
        ("Acceso a varios modelos por una sola API. ", True),
        ("El caso necesita tres familias distintas —lenguaje, difusión y embeddings— de "
         "tres proveedores distintos. Bedrock las expone con una única interfaz y un "
         "único mecanismo de credenciales.", False),
    ])

    parrafo(doc,
            "La contrapartida se asume de forma consciente: dependencia de un "
            "proveedor, menor control sobre las versiones de los modelos y coste "
            "marginal creciente si el volumen se disparara. El punto en el que "
            "convendría reconsiderar la decisión es identificable: un volumen sostenido "
            "y predecible que amortice una GPU reservada, o un requisito regulatorio "
            "que impida que el contenido salga de la propia infraestructura.")

    doc.add_heading("La decisión estructural: dónde se simula", level=2)

    parrafo(doc,
            "Al no disponer de cuenta AWS, había que decidir en qué punto sustituir la "
            "llamada real. La opción evidente —una bifurcación en la lógica de negocio— "
            "se descartó porque deja el código de integración sin ejercitar: se "
            "escribe, pero nunca se comprueba que su forma sea correcta.")

    parrafo(doc,
            "La solución adoptada sustituye el objeto de transporte, no la lógica. "
            "MockBedrockRuntime expone la misma operación invoke_model que el cliente "
            "de boto3, recibe exactamente el mismo cuerpo JSON que se enviaría a AWS y "
            "devuelve una respuesta con la misma estructura de sobre. En consecuencia, "
            "todo el código de la aplicación es código de producción: construye los "
            "payloads reales, los serializa igual y parsea las respuestas por el mismo "
            "camino.")

    caja_destacada(
        doc,
        "Consecuencia verificable",
        "El sustituto valida cada payload que recibe contra el contrato documentado de "
        "cada modelo y rechaza los mal formados igual que haría AWS: anthropic_version "
        "incorrecto, temperature fuera de [0, 1], cfg_scale fuera de rango, "
        "text_prompts vacío. Una petición mal construida falla en desarrollo en lugar "
        "de devolver una respuesta plausible que oculte el error. La suite "
        "tests/test_contratos_bedrock.py comprueba esas validaciones.",
        "E8F0F8",
    )


def seccion_33(doc):
    doc.add_heading("3.3. Modelos y parámetros de inferencia", level=1)

    parrafo(doc,
            "Cada tarea usa el modelo y los parámetros que le corresponden por su "
            "naturaleza, siguiendo la regla de la unidad: rigor y consistencia exigen "
            "temperatura baja; creatividad y diversidad, temperatura alta. La tabla "
            "siguiente se genera automáticamente a partir de src/bedrock/models.py, de "
            "modo que refleja exactamente lo que la aplicación ejecuta.")

    doc.add_heading("Operaciones de texto (Claude)", level=2)

    filas = []
    orden = [TextTask.CORREGIR, TextTask.RESUMIR, TextTask.EXPANDIR, TextTask.VARIAR]
    for tarea in orden:
        p = TEXT_PROFILES[tarea]
        filas.append([
            tarea.etiqueta,
            p.model_id.replace("anthropic.", ""),
            f"{p.temperature}",
            f"{p.top_p}",
            f"{p.max_tokens}",
        ])
    tabla(doc, ["Tarea", "Modelo", "Temp.", "Top-P", "Máx. tokens"], filas,
          anchos=[4.4, 4.6, 1.9, 1.9, 3.2])

    parrafo(doc, "Justificación de cada elección:", negrita=True)
    for tarea in orden:
        p = TEXT_PROFILES[tarea]
        parrafo_rico(doc, [(f"{tarea.etiqueta}. ", True), (p.justificacion, False)],
                     sangria=0.5, espacio_despues=6)

    doc.add_heading("Por qué estos dos modelos y no los más recientes", level=2)

    parrafo(doc,
            "Se emplean Claude Haiku 4.5 para las tareas deterministas y de alto "
            "volumen, y Claude Sonnet 4.6 para las creativas. La razón de no usar los "
            "modelos más recientes de la familia Opus es técnica y merece explicarse: "
            "esos modelos sustituyen el muestreo explícito por razonamiento adaptativo "
            "y un parámetro de esfuerzo, y rechazan el parámetro temperature. Adoptarlos "
            "habría hecho imposible demostrar el control de inferencia que es el objeto "
            "de estudio de esta unidad.")

    parrafo(doc,
            "Merece la pena registrar la implicación de fondo: el control de la "
            "generación se está desplazando del muestreo estadístico —temperatura, "
            "Top-P— hacia el control del proceso de razonamiento. Un diseño hecho hoy "
            "sobre temperatura debe prever esa migración.")

    doc.add_heading("Generación de imágenes (Stable Diffusion XL)", level=2)

    parrafo(doc,
            "En imagen los parámetros no son temperatura y Top-P, sino la escala de "
            "guía (cfg_scale), los pasos de difusión (steps) y la semilla (seed). Cada "
            "estilo combina dos palancas: el preset nativo del modelo y un sufijo de "
            "prompt, porque los presets no cubren todos los estilos que pide el "
            "enunciado —no existe un preset de pintura al óleo—.")

    filas_img = [
        [e.etiqueta, e.style_preset, f"{e.cfg_scale}", f"{e.steps}"]
        for e in IMAGE_STYLES.values()
    ]
    tabla(doc, ["Estilo", "style_preset", "cfg_scale", "steps"], filas_img,
          anchos=[5.0, 4.6, 3.2, 3.2])

    for estilo in IMAGE_STYLES.values():
        parrafo_rico(doc, [(f"{estilo.etiqueta}. ", True), (estilo.justificacion, False)],
                     sangria=0.5, espacio_despues=6)

    caja_destacada(
        doc,
        "La semilla es el parámetro que hace auditable la generación",
        "Con la misma semilla y el mismo prompt, Stable Diffusion produce la misma "
        "imagen. Eso convierte la generación en un proceso repetible: el diseñador "
        "puede iterar sobre una propuesta concreta en lugar de volver a tirar los "
        "dados, y la agencia puede acreditar cómo se produjo una imagen publicada. Por "
        "eso la semilla se guarda junto a cada pieza de la galería y no se trata como "
        "un detalle interno.",
        "FFF4E5",
    )

    doc.add_heading("Embeddings (Amazon Titan)", level=2)
    parrafo_rico(doc, [
        ("Modelo: ", True), (f"{TITAN_EMBEDDINGS}. ", False),
        ("Se solicitan vectores de 1024 dimensiones normalizados. La normalización se "
         "pide al modelo porque permite calcular la similitud coseno como un simple "
         "producto escalar, sin dividir por las normas en cada consulta.", False),
    ])

    doc.add_heading("Control de coste y de longitud", level=2)
    parrafo(doc,
            "max_tokens se fija por tarea y no de forma global: 1024 para resumir "
            "—porque un resumen más largo ha dejado de ser un resumen— y 4096 para "
            "expandir. Es simultáneamente un control de calidad y un control de gasto, "
            "ya que en Bedrock se paga por token generado. En producción se añadiría "
            "además un presupuesto mensual con alertas, y caché de prompts para los "
            "system prompts, que son idénticos en todas las llamadas de una misma "
            "tarea.")


def seccion_34(doc):
    doc.add_heading("3.4. Instrucciones del sistema (System Prompt)", level=1)

    parrafo(doc,
            "El system prompt se construye por composición, no como cuatro textos "
            "independientes: un tronco común e invariable —rol, restricciones y reglas "
            "de prioridad— más un bloque específico por operación con el objetivo y el "
            "formato de salida. La razón es de seguridad: si cada operación repitiera "
            "sus propias reglas, acabarían divergiendo y la defensa frente a inyección "
            "tendría agujeros distintos según el botón que pulsara el usuario.")

    doc.add_heading("Los cuatro pilares", level=2)

    tabla(doc,
          ["Pilar", "Contenido en Aurora Studio"],
          [
              ["Rol / persona",
               "Asistente editorial interno de una agencia de marketing. Editor "
               "profesional de marca: preciso, sobrio y directo. Prioriza la claridad "
               "sobre el lucimiento."],
              ["Objetivo",
               "Específico por operación: resumir al 30–40 %, expandir sin introducir "
               "hechos nuevos, corregir sin reescribir por preferencia, generar tres "
               "variaciones genuinamente distintas."],
              ["Restricciones",
               "Siete reglas inquebrantables: no inventar datos, avisar cuando falte "
               "información, respetar el idioma, conservar el significado, no emitir "
               "afirmaciones sensibles sin respaldo, no reproducir marcas ajenas, no "
               "producir estereotipos. Más las reglas de prioridad frente a inyección."],
              ["Formato de salida",
               "Definido por operación: texto plano sin preámbulos para resumir; "
               "párrafos con máximo tres secciones para expandir; texto corregido más "
               "una sección CAMBIOS para corregir; tres bloques etiquetados en Markdown "
               "para variar."],
          ],
          anchos=[3.4, 12.6])

    doc.add_heading("System prompt completo (operación: corregir)", level=2)

    parrafo(doc,
            "Se reproduce íntegro el prompt que la aplicación envía, extraído "
            "directamente del código. El bloque de contexto de marca que añade el RAG "
            "se muestra en la sección 3.5.")

    prompt = build_system_prompt(TextTask.CORREGIR)
    bloque_codigo(doc, prompt, tamano=7.8)

    doc.add_heading("La regla que hace el trabajo de seguridad", level=2)

    parrafo(doc,
            "El apartado «Prioridad de instrucciones» es el que convierte la "
            "delimitación en una defensa real. Declara explícitamente que lo que llega "
            "entre las etiquetas de contenido es dato y nunca instrucción, enumera los "
            "patrones de ataque más habituales para que el modelo los reconozca como "
            "texto a editar, y establece que las reglas de sistema prevalecen sobre "
            "cualquier cosa que aparezca en el contenido del usuario.")

    parrafo(doc,
            "Por sí sola no basta: un usuario podría cerrar la etiqueta e intentar "
            "escribir fuera del bloque delimitado. Esa vía se cierra en código, no en "
            "el prompt, y se explica en la sección 3.6.")


def seccion_35(doc):
    doc.add_heading("3.5. RAG y memoria", level=1)

    caja_destacada(
        doc,
        "La distinción, en una línea",
        "RAG aporta conocimiento verificable; la memoria aporta continuidad de trabajo. "
        "Aurora Studio implementa ambos, y los mantiene separados a propósito.",
        "E8F4EC",
    )

    doc.add_heading("RAG: sí, y por qué", level=2)

    parrafo(doc,
            "La aplicación necesita que Claude conozca la guía de estilo interna: "
            "paleta, terminología de producto, prohibiciones de estilo y política "
            "legal. Ese conocimiento no puede vivir en el system prompt por tres "
            "razones: ocuparía miles de tokens en cada llamada, quedaría desactualizado "
            "cada vez que el departamento legal cambiara una norma, y no permitiría "
            "citar la fuente de una recomendación concreta.")

    parrafo(doc,
            "Es el caso de uso canónico de RAG: recuperar de un corpus propio los "
            "fragmentos relevantes para la consulta y entregárselos al modelo antes de "
            "que responda.")

    doc.add_heading("Cómo está construido el índice", level=2)

    bloque_codigo(doc,
                  "4 documentos .md de la guía de marca\n"
                  "        |\n"
                  "        v  troceado por secciones (encabezados de nivel 2)\n"
                  "14 fragmentos\n"
                  "        |\n"
                  "        v  Amazon Titan Embeddings · 1024 dimensiones\n"
                  "matriz 14 x 1024\n"
                  "        |\n"
                  "consulta --> embedding --> similitud coseno --> top-3 con umbral",
                  tamano=9)

    parrafo(doc,
            "Tres decisiones de diseño merecen justificarse:")

    vineta_rica(doc, [
        ("Troceado por secciones, no por ventana de caracteres. ", True),
        ("La guía ya viene estructurada por temas. Partirla cada N caracteres rompería "
         "las tablas y separaría una prohibición de su contexto. Cuando el documento "
         "tiene estructura propia, el troceado debe respetarla.", False),
    ])
    vineta_rica(doc, [
        ("Índice en memoria, no base vectorial gestionada. ", True),
        ("El corpus son catorce fragmentos. Introducir OpenSearch Serverless o pgvector "
         "aquí sería sobreingeniería. El punto de corte a partir del cual sí "
         "compensaría está en el orden de las decenas de miles de fragmentos, o cuando "
         "el índice deba compartirse entre procesos y sobrevivir a los reinicios.", False),
    ])
    vineta_rica(doc, [
        ("Umbral de similitud calibrado empíricamente. ", True),
        ("Se midió la puntuación del mejor fragmento en dos poblaciones: ocho consultas "
         "legítimas sobre la guía y cuatro consultas de ruido ajenas al dominio. El "
         "ruido no superó 0,052 y la consulta legítima peor puntuada alcanzó 0,120. El "
         "umbral se fijó en 0,09, en el hueco entre ambas. El script "
         "tests/test_rag_umbral.py reproduce la medición.", False),
    ])

    parrafo(doc,
            "El umbral evita el fallo más común de un RAG mal calibrado: entregar "
            "siempre k fragmentos aunque ninguno venga a cuento. Inyectar ruido es peor "
            "que no inyectar nada, porque empuja al modelo a forzar conexiones que no "
            "existen.")

    doc.add_heading("Memoria: sí, y es otra cosa", level=2)

    parrafo(doc,
            "La memoria de Aurora Studio es el historial de la pieza en curso: el texto "
            "que se está editando, sus versiones anteriores, quién hizo cada cambio con "
            "qué modelo y parámetros, y los comentarios asociados. Aporta continuidad "
            "dentro de la sesión de trabajo.")

    parrafo(doc,
            "Está implementada como un registro de solo-anexado: las versiones no se "
            "sobrescriben nunca, y restaurar una versión anterior apila una versión "
            "nueva en lugar de borrar. Es más caro en memoria y es lo correcto, porque "
            "en un flujo de aprobación poder demostrar qué se aprobó y cuándo es un "
            "requisito de trazabilidad, no una comodidad.")

    doc.add_heading("Por qué no se confunden", level=2)

    tabla(doc,
          ["", "RAG", "Memoria"],
          [
              ["Qué aporta", "Conocimiento verificable y externo",
               "Continuidad del trabajo en curso"],
              ["De dónde viene", "Guía de estilo de marca (documentos propios)",
               "Acciones del usuario en esta sesión"],
              ["Cómo se recupera", "Similitud semántica sobre embeddings",
               "Acceso directo al historial de la pieza"],
              ["Cuándo cambia", "Cuando legal o marca actualizan la guía",
               "En cada operación que el usuario guarda"],
              ["Dónde vive", "Índice vectorial (src/services/rag_service.py)",
               "Estado de sesión (src/domain/versioning.py)"],
          ],
          anchos=[3.2, 6.4, 6.4])

    parrafo(doc,
            "El error que se ha evitado deliberadamente es usar la memoria "
            "conversacional como sustituto del RAG, es decir, pegar la guía de estilo "
            "en el primer mensaje de la conversación y confiar en que el modelo la "
            "recuerde. Eso consume contexto en cada turno, degrada la atención sobre lo "
            "importante y no permite saber qué fragmento concreto respaldó una "
            "recomendación.")


def seccion_36(doc):
    doc.add_heading("3.6. Ética y seguridad", level=1)

    parrafo(doc,
            "Los cuatro frentes se abordan con controles implementados en el flujo de "
            "ejecución, no solo declarados. Se documenta también, en cada uno, hasta "
            "dónde llega el control y qué haría falta en producción.")

    doc.add_heading("Moderación", level=2)

    parrafo(doc,
            "Se modera la entrada y también la salida. Moderar solo la entrada es "
            "insuficiente: un prompt inocuo puede producir una salida problemática. El "
            "orden de los pasos está pensado para que el control barato se ejecute "
            "antes que el caro:")

    bloque_codigo(doc,
                  "1. Moderar la entrada        <- si bloquea, no se gasta en el modelo\n"
                  "2. Sanear anti-inyeccion     <- antes de construir el prompt\n"
                  "3. Recuperar contexto RAG\n"
                  "4. Construir el prompt\n"
                  "5. Invocar el modelo\n"
                  "6. Moderar la salida         <- el riesgo no esta solo en la entrada",
                  tamano=9)

    parrafo(doc,
            "La moderación implementada es por listas y reglas: detecta lo evidente y "
            "deja pasar lo sutil. Se declara así de forma explícita porque presentarla "
            "como una solución completa sería precisamente el error que el caso pide "
            "evitar. En una implantación real sería la primera de tres capas: esta, "
            "más Amazon Bedrock Guardrails con políticas gestionadas y detección de "
            "información personal, más los filtros nativos de los propios modelos "
            "—Stable Diffusion ya devuelve finishReason CONTENT_FILTERED cuando "
            "rechaza una generación, y la aplicación lo trata—.")

    doc.add_heading("Inyección de prompt", level=2)

    parrafo(doc,
            "El modelo de amenaza asumido es un usuario autenticado de la herramienta: "
            "un redactor, o alguien que pega un texto de origen externo sin revisarlo. "
            "Su objetivo puede ser extraer el system prompt, anular las reglas de marca "
            "o usar la herramienta fuera de política.")

    parrafo(doc, "La defensa tiene tres capas, y solo una es determinista:")

    tabla(doc,
          ["Capa", "Qué hace", "Garantía"],
          [
              ["Neutralización de delimitadores",
               "Sustituye cualquier aparición de las etiquetas reservadas en el texto "
               "del usuario, de modo que no pueda cerrar el bloque de contenido.",
               "Determinista"],
              ["Detección de patrones",
               "Identifica frases con forma de instrucción («ignora las instrucciones "
               "anteriores», «revela tu prompt») y avisa al usuario y al registro.",
               "Heurística"],
              ["Prioridad declarada en el prompt",
               "El system prompt establece que lo delimitado es dato y que las reglas "
               "de sistema prevalecen.",
               "Depende del modelo"],
          ],
          anchos=[4.4, 8.6, 3.0])

    caja_destacada(
        doc,
        "Una decisión de producto, no solo técnica",
        "La entrada sospechosa no se rechaza: se neutraliza y se avisa. Bloquearla "
        "sería la respuesta equivocada, porque un redactor puede legítimamente querer "
        "editar un artículo sobre inyección de prompt, y la herramienta debe "
        "permitirlo. Lo que no debe permitir es que ese texto cambie el comportamiento "
        "del modelo.",
        "FFF4E5",
    )

    parrafo(doc,
            "Como control de última línea, la moderación de salida comprueba que la "
            "respuesta no contenga fragmentos del propio system prompt. Si los "
            "contuviera, sería la señal de una inyección con éxito y la salida se "
            "bloquea antes de mostrarse.")

    doc.add_heading("Sesgo", level=2)

    parrafo(doc,
            "El sesgo se aborda en tres puntos. En el system prompt, con una regla "
            "inquebrantable que prohíbe estereotipos y obliga a señalar los que "
            "aparezcan en el texto de entrada. En los prompts de imagen, con un "
            "detector de descriptores que fijan atributos personales sin necesidad "
            "funcional: un rol profesional sin más contexto hará que el modelo "
            "reproduzca el estereotipo dominante de su conjunto de entrenamiento. Y en "
            "la salida, revisando el texto generado con los mismos marcadores.")

    parrafo(doc,
            "El detector avisa, no bloquea, y la razón es la misma que antes: el "
            "objetivo es que la persona decida con la información delante, no que la "
            "herramienta decida por ella. Un filtro de sesgo que bloquea acaba "
            "desactivado.")

    parrafo(doc,
            "Límite declarado: esto detecta sesgo léxico en el prompt, no sesgo en la "
            "imagen resultante. Evaluar lo segundo exigiría generar lotes y analizar "
            "la distribución demográfica de las salidas, que es un trabajo de auditoría "
            "periódica, no un control en línea.")

    doc.add_heading("Privacidad y derechos de autor", level=2)

    vineta_rica(doc, [
        ("Marcas registradas. ", True),
        ("El prompt de imagen se contrasta con una lista de marcas protegidas. Generar "
         "piezas que evoquen marcas ajenas expone a la agencia a una reclamación por "
         "infracción, así que aquí sí se bloquea.", False),
    ])
    vineta_rica(doc, [
        ("Estilos de autor. ", True),
        ("Se bloquean las peticiones que invocan el estilo de un artista vivo o de un "
         "estudio con obra protegida. El estilo debe describirse por sus atributos "
         "visuales, no por su autor. Es el vector de riesgo más habitual en generación "
         "de imágenes para marketing.", False),
    ])
    vineta_rica(doc, [
        ("Procedencia. ", True),
        ("Cada imagen conserva su prompt, su estilo, su semilla, su autor y la marca "
         "explícita de haber sido generada con IA. Es lo que permite acreditar el "
         "proceso si un tercero lo cuestiona, y es un requisito que la propia guía de "
         "marca impone.", False),
    ])
    vineta_rica(doc, [
        ("Datos personales. ", True),
        ("La política prohíbe enviar datos personales de clientes a los modelos. En "
         "producción esto se reforzaría con la detección de información personal de "
         "Bedrock Guardrails, que puede redactar automáticamente los datos antes de la "
         "inferencia.", False),
    ])
    vineta_rica(doc, [
        ("Cifrado. ", True),
        ("En producción: cifrado en reposo con claves gestionadas por el cliente en KMS "
         "para el bucket de imágenes y el historial, y TLS en tránsito, que Bedrock ya "
         "impone. En esta entrega el almacenamiento es de sesión y no persiste, por lo "
         "que el cifrado en reposo no aplica todavía; se documenta como requisito del "
         "paso a producción, no como algo resuelto.", False),
    ])

    parrafo(doc,
            "Sobre la titularidad de lo generado: es un terreno jurídicamente "
            "inestable que varía por jurisdicción. La política adoptada es de "
            "prudencia —no registrar las imágenes generadas como obra propia, no usarlas "
            "como elemento central de una identidad de marca registrable, y documentar "
            "siempre el prompt y la semilla—. No es una respuesta jurídica; es una "
            "posición defendible mientras no la haya.")


def seccion_4(doc):
    doc.add_heading("4. Vía A · La aplicación entregada", level=1)

    parrafo(doc,
            "Aurora Studio es una aplicación Streamlit con cuatro pantallas. Cubre el "
            "alcance mínimo que exige la guía y la totalidad del alcance ampliado.")

    doc.add_heading("Alcance cubierto", level=2)

    tabla(doc,
          ["Requisito de la guía", "Estado", "Dónde se ve"],
          [
              ["Campo de texto que genera imagen con Stable Diffusion vía Bedrock",
               "Completo", "Pantalla «Generación de imágenes»"],
              ["Función que mejora texto con Claude (resumir y corregir como mínimo)",
               "Completo · 4 operaciones", "Pantalla «Edición de contenido»"],
              ["Ambos flujos llaman a Bedrock y muestran el resultado",
               "Completo · modo simulado", "Panel «Detalle técnico de la llamada»"],
              ["Selección de estilos y galería con descarga",
               "Completo · 4 estilos", "Pantallas «Generación» y «Galería»"],
              ["Historial y control de versiones del texto",
               "Completo · con comparación visual", "Pantalla «Colaboración»"],
              ["Roles, permisos y comentarios",
               "Completo · 3 roles, 6 permisos", "Pantalla «Colaboración»"],
          ],
          anchos=[7.4, 4.2, 4.4])

    doc.add_heading("Cómo ejecutarla", level=2)
    bloque_codigo(doc,
                  "pip install -r requirements.txt\n"
                  "python -m streamlit run app.py\n"
                  "\n"
                  "# Para llamar a AWS de verdad, sin tocar el codigo:\n"
                  "#   copiar .env.example a .env\n"
                  "#   BEDROCK_BACKEND=aws\n"
                  "#   AWS_REGION=us-east-1\n"
                  "#   aws configure",
                  tamano=9)
    parrafo(doc,
            "Se invoca Streamlit como módulo de Python y no con el comando "
            "«streamlit» a secas porque en muchas instalaciones de Windows la "
            "carpeta Scripts del intérprete no está en el PATH, y el comando "
            "directo falla aunque el paquete esté bien instalado. Invocarlo como "
            "módulo usa el mismo intérprete con el que se instalaron las "
            "dependencias.", tamano=10, color=GRIS)

    doc.add_heading("Evidencia de la integración", level=2)

    parrafo(doc,
            "Como la aplicación no llama a AWS, la evidencia de que la integración es "
            "correcta se aporta por tres vías distintas:")

    vineta_rica(doc, [
        ("El payload visible en la interfaz. ", True),
        ("Cada operación muestra, en el panel «Detalle técnico de la llamada», el JSON "
         "exacto que se envía a Bedrock. Es el mismo en modo simulado y en modo real.", False),
    ])
    vineta_rica(doc, [
        ("Validación de contrato en el sustituto. ", True),
        ("MockBedrockRuntime rechaza los payloads mal formados con los mismos criterios "
         "que AWS.", False),
    ])
    vineta_rica(doc, [
        ("Pruebas automatizadas. ", True),
        ("tests/test_contratos_bedrock.py comprueba la forma de las peticiones de los "
         "tres modelos y el comportamiento de los controles de seguridad; "
         "tests/test_rag_umbral.py reproduce la calibración del RAG.", False),
    ])


def seccion_5(doc):
    doc.add_heading("5. Limitaciones declaradas", level=1)

    parrafo(doc,
            "Enumerarlas es parte del trabajo: un diseño que no reconoce sus límites no "
            "es un diseño terminado.")

    tabla(doc,
          ["Limitación", "Alcance real", "Qué haría falta"],
          [
              ["No se ejecuta contra AWS",
               "Ninguna respuesta procede de los modelos reales.",
               "Cuenta AWS con acceso concedido a los tres modelos y "
               "BEDROCK_BACKEND=aws."],
              ["Calidad del sustituto de embeddings",
               "Es un recuperador léxico, no semántico. Recall@3 medido: 5 de 8 "
               "consultas. Falla en las que exigen entender el significado, no "
               "emparejar palabras.",
               "Titan real. Se espera que suba sin tocar código de aplicación."],
              ["Moderación por listas",
               "Detecta lo evidente; no es un clasificador.",
               "Bedrock Guardrails como segunda capa."],
              ["Autorización sin autenticación",
               "El selector de usuario simula una sesión iniciada. No hay contraseñas "
               "ni verificación de identidad.",
               "Amazon Cognito o el SSO corporativo. El módulo de permisos ya está "
               "separado y no habría que reescribirlo."],
              ["Persistencia de sesión",
               "Galería e historial viven en memoria y se pierden al reiniciar.",
               "DynamoDB para el historial y S3 para las imágenes, con cifrado KMS."],
              ["Concurrencia real",
               "La colaboración es multiusuario en el modelo de roles, pero un solo "
               "proceso de Streamlit.",
               "Backend con estado compartido y bloqueo optimista por versión."],
          ],
          anchos=[4.0, 6.6, 5.4])


def seccion_6(doc):
    doc.add_heading("6. Autoevaluación frente a la guía", level=1)

    filas = [
        ["Problema e historias de usuario (diseñador, redactor, aprobador)", "Sí", "3.1"],
        ["Diagrama con interfaz, wrapper, Bedrock y almacenamiento", "Sí", "3.2"],
        ["Justificación de API gestionada frente a open source propio", "Sí", "3.2"],
        ["Modelo y parámetros por tarea, con su porqué", "Sí", "3.3"],
        ["System prompt con los cuatro pilares", "Sí", "3.4"],
        ["Uso justificado —y no confundido— de RAG y memoria", "Sí", "3.5"],
        ["Moderación, inyección de prompt, sesgo, privacidad y copyright", "Sí", "3.6"],
        ["La app genera imagen y mejora texto llamando a Bedrock", "Sí · simulado", "4"],
        ["Código con README que explica cómo ejecutarlo", "Sí", "README.md"],
        ["Demostración con capturas comentadas", "Guion preparado", "docs/GUION_DEMO.md"],
        ["Fuentes citadas", "Sí", "7"],
    ]
    tabla(doc, ["Punto de la guía", "Cumplido", "Sección"], filas,
          anchos=[10.6, 3.0, 2.4])


def seccion_7(doc):
    doc.add_heading("7. Fuentes y herramientas", level=1)

    parrafo(doc, "Documentación oficial consultada:", negrita=True)
    for fuente in [
        "Amazon Bedrock — Guía del usuario y referencia de la API de inferencia "
        "(estructura de invoke_model, formatos de petición y respuesta por proveedor).",
        "Anthropic Claude en Amazon Bedrock — parámetros de inferencia y formato de la "
        "Messages API (anthropic_version, system, messages, temperature, top_p).",
        "Stability AI Stable Diffusion XL en Bedrock — contrato de text_prompts, "
        "cfg_scale, steps, seed y style_preset.",
        "Amazon Titan Embeddings — parámetros inputText, dimensions y normalize.",
        "Amazon Bedrock Guardrails — políticas de contenido y detección de información "
        "personal (citado como trabajo futuro, no implementado).",
        "Documentación de Streamlit — gestión del estado de sesión y ciclo de "
        "reejecución.",
    ]:
        vineta(doc, fuente)

    parrafo(doc, "Herramientas empleadas:", negrita=True)
    for herramienta in [
        "Python 3.14 · Streamlit 1.63 · boto3 · NumPy · Pillow · Matplotlib · python-docx",
        "Claude Code (Anthropic) como asistente de desarrollo durante la implementación.",
    ]:
        vineta(doc, herramienta)

    parrafo(doc,
            "Los materiales de la asignatura —enunciado del caso y guía del trabajo "
            "práctico de la Unidad 3— son la base de los requisitos recogidos en este "
            "documento.", tamano=10, color=GRIS)


# ---------------------------------------------------------------------------


def construir() -> Path:
    doc = Document()

    seccion = doc.sections[0]
    seccion.top_margin = Cm(2.2)
    seccion.bottom_margin = Cm(2.2)
    seccion.left_margin = Cm(2.4)
    seccion.right_margin = Cm(2.4)

    configurar_estilos(doc)
    pie_de_pagina(doc)

    portada(doc)
    seccion_indice(doc)
    seccion_1(doc)
    doc.add_page_break()

    doc.add_heading("Tronco común", level=1)
    parrafo(doc,
            "Las seis secciones que la guía exige por igual en ambas vías. Se conserva "
            "deliberadamente su numeración original (3.1 a 3.6) para que el contraste "
            "punto por punto con la guía del trabajo sea inmediato.",
            cursiva=True, color=GRIS)

    seccion_31(doc)
    seccion_32(doc)
    doc.add_page_break()
    seccion_33(doc)
    doc.add_page_break()
    seccion_34(doc)
    doc.add_page_break()
    seccion_35(doc)
    doc.add_page_break()
    seccion_36(doc)
    doc.add_page_break()
    seccion_4(doc)
    seccion_5(doc)
    doc.add_page_break()
    seccion_6(doc)
    seccion_7(doc)

    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    destino = DOCS_DIR / "Memoria_CasoPractico_U3_Jose_Ruber_Moncayo_Navia.docx"
    doc.save(destino)
    return destino


if __name__ == "__main__":
    ruta = construir()
    print(f"Memoria generada: {ruta}")
